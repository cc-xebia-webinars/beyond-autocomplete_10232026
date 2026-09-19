"""Build the webinar slide deck from the Xebia template.

This script is the only way the deck changes. The PPTX in `docs/` is generated from
the SLIDES list below and should never be edited by hand: change the data here,
rebuild, regenerate the PDF, and then update the matching heading in
PRESENTER_SCRIPT.md.

Run from the repository root:

    uv run --with python-pptx python instructor/build_deck.py
    uv run --with python-pptx python instructor/build_deck.py --inventory
    uv run --with python-pptx python instructor/build_deck.py --check-script

Then convert to PDF with LibreOffice, as described in the root CLAUDE.md.
"""

from __future__ import annotations

import argparse
import copy
import io
import re
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "instructor" / "Xebia PPT General Template.pptx"
OUTPUT = ROOT / "docs" / "agentic-coding-webinar-slides.pptx"
SCRIPT = ROOT / "instructor" / "PRESENTER_SCRIPT.md"

# Diagrams are drawn with the diagram-builder MCP server. Each .svg in this folder is
# the editable source, and the .png next to it is the export the deck embeds.
DIAGRAMS = ROOT / "instructor" / "diagrams"

# Layout indexes in the Xebia template, named so the slide list below reads clearly.
COVER = 0          # Main-cover_light
BASIC = 6          # Content_Basic
THREE_COL = 8      # Content_3 Columns_Comparison_2
TWO_COL = 9        # Content_2 Columns
BREAK = 24         # 1_Photo_background
END = 27           # End-cover_light

# The template's own body default is 12 pt, which is unreadable on a shared screen,
# so every run gets an explicit size instead of inheriting one.
BULLET_PT = 20
CODE_PT = 18
COLUMN_CODE_PT = 16
COLUMN_BODY_PT = 16
COLUMN_HEADING_PT = 16
SUBTITLE_PT = 16
MONO = "Consolas"

# Density rules from CLAUDE.md, checked against this data before anything is built.
MAX_SLIDES = 20
MAX_BULLETS = 6
MAX_WORDS = 11     # "under a dozen words"
MAX_CODE_LINES = 12

# Column placeholders aren't in left-to-right index order, so each column is named by
# its (heading idx, body idx) pair, measured from the template.
THREE_COL_SLOTS = [(11, 1), (17, 16), (14, 13)]    # left, middle, right
TWO_COL_SLOTS = [(11, 1), (14, 13)]                # left, right

# The two-column layout reserves a band for icon pictures above each column. We don't
# use icons, and once those frames are removed the columns sit alone in the bottom half
# of the slide. Moving them up to the three-column layout's positions closes the gap.
TWO_COL_HEADING_TOP_IN = 2.16
TWO_COL_BODY_TOP_IN = 2.84
TWO_COL_BODY_HEIGHT_IN = 4.06

# Diagram slides fill the body area of the basic layout, under the title.
DIAGRAM_LEFT_IN = 0.61
DIAGRAM_TOP_IN = 1.85
DIAGRAM_MAX_HEIGHT_IN = 5.0
DIAGRAM_MAX_WIDTH_IN = 12.1


# ---------------------------------------------------------------------------- slides
#
# Text wrapped in backticks is set in a monospace font, so method names, file paths
# and code read as code. Everything else is set in the template's Arial.

SLIDES = [
    {
        "kind": "Title",
        "layout": COVER,
        # The line break keeps the title clear of the cover's artwork.
        "title": "Agentic Coding\nBeyond Autocomplete",
        "subtitle": ("Real Work with Copilot Agent Mode and Claude Code\n"
                     "Eric Greene · Xebia\neric.greene@xebia.com"),
        "notes": (
            "Welcome and introductions. Ninety minutes: about twenty of slides and seventy of "
            "live demos in VS Code, with GitHub Copilot's agent mode and the Claude Code "
            "extension. Questions are welcome at any time, in the chat or out loud. Everything "
            "is in the webinar repo. Set the frame: autocomplete is table stakes, and the "
            "interesting question now is how to get real work done with agents."
        ),
    },
    {
        "kind": "Agenda",
        "layout": BASIC,
        "title": "Agenda",
        "bullets": [
            "From inline suggestions to autonomous agents",
            "How Copilot agent mode works",
            "How Claude Code works",
            "Patterns: refactors, tests, debugging, and new features",
            "Keeping agents on track: prompts, intervention, and review",
            "Team adoption: repo hygiene, context, and workflow",
        ],
        "notes": (
            "Six topics, matching the README agenda, and four demos on one codebase, Roombook, "
            "a small API for booking meeting rooms in Amsterdam and New York. Topics one and two "
            "lead into Demo 1 with Copilot, topic three into Demo 2 with Claude Code, topic four "
            "into Demo 3, and topics five and six into Demo 4. Each demo starts where the last "
            "one stopped."
        ),
    },
    {
        "kind": "Content",
        "layout": THREE_COL,
        "title": "From suggestions to agents",
        "columns": [
            ("Autocomplete", [
                "Finishes the line you're typing",
                "You drive every keystroke",
            ]),
            ("Chat", [
                "Answers questions, drafts code",
                "You copy, paste, and run it",
            ]),
            ("Agent", [
                "Plans, edits, runs, and iterates",
                "You set the goal and review",
            ]),
        ],
        "notes": (
            "Three generations of AI coding help, and all three are still useful. Autocomplete "
            "works inside the line or function you're writing. Chat answers questions, but you're "
            "the one who moves the code into place and runs it. An agent takes a goal and works "
            "toward it on its own: it reads the codebase, edits files, runs commands, and checks "
            "the results. The bottom row is the real shift: your role moves from typing to "
            "directing and reviewing."
        ),
    },
    {
        "kind": "Diagram",
        "layout": BASIC,
        "title": "The agent loop",
        "image": "agent-loop.png",
        # Trims the diagram page's empty top and bottom margins, as fractions.
        "crop": (0.27, 0.2),
        "notes": (
            "Every coding agent runs some version of this loop. It gathers context, plans, acts "
            "by editing files and running commands, and verifies with tests and errors, then goes "
            "round again until the checks pass. The orange row is where you stay in the loop: "
            "the context you give it, the plan you review, the commands you approve, and the diff "
            "and tests you review at the end. Both tools today work this way."
        ),
    },
    {
        "kind": "Content",
        "layout": TWO_COL,
        "title": "What changes, and what breaks",
        "columns": [
            ("What changes", [
                "Scope grows from a line to a codebase",
                "The agent runs commands, not just you",
                "Your job shifts to direction and review",
            ]),
            ("What breaks", [
                "Review load grows with every file touched",
                "Confident changes built on missing context",
                "Tests quietly edited until they pass",
            ]),
        ],
        "notes": (
            "What changes is scope and autonomy. An agent can touch ten files and run your test "
            "suite in the time it takes to read one suggestion. What breaks is the old habit of "
            "glancing at a suggestion and pressing Tab. You now review more code, the agent will "
            "sound just as sure when it's missing context, and an agent that's told to make the "
            "tests pass may change the tests. The rest of the session is about handling those three."
        ),
    },
    {
        "kind": "Content",
        "layout": BASIC,
        "title": "How Copilot agent mode works",
        "bullets": [
            "Ask, Plan, and Agent, in the Copilot harness",
            "Searches the workspace and reads files itself",
            "Edits several files, reviewed as one diff",
            "Runs commands after you select Allow",
            "Reads test output and errors, then tries again",
            "Checkpoints roll the workspace back to any request",
        ],
        "notes": (
            "We'll use the Copilot harness, VS Code's default since 1.143, which replaces the "
            "older Local agent. Ask answers questions, Plan writes a plan you can comment on, and "
            "Agent makes changes. It finds the files it needs without being told, edits across the "
            "codebase, and runs commands once you allow them, then reads the output and iterates. "
            "You review its edits as one diff, and Restore Checkpoint throws away a run."
        ),
    },
    {
        "kind": "Demo",
        "layout": BREAK,
        "title": "Demo 1: From autocomplete to agent",
        "subtitle": "demos/01-copilot-agent-mode · about 17 minutes",
        "notes": (
            "Run the tests and the API, read create_booking. Ghost text for has_equipment, then a "
            "next edit suggestion for a rename that breaks a caller in another file. Agent role: "
            "GET /rooms/available with tests. Watch the context gathering, Allow on uv run pytest, "
            "and the iteration if the new route returns 422. Review the changes as one diff, "
            "restore a checkpoint, and run the tests."
        ),
    },
    {
        "kind": "Content",
        "layout": BASIC,
        "title": "How Claude Code works",
        "bullets": [
            "Gathers context with search and read tools",
            "Loads CLAUDE.md into every session",
            "Edits files and runs commands as tool calls",
            "Every tool call passes a permission check",
            "Plan mode thinks before it changes anything",
            "Rewind undoes edits you don't want",
        ],
        "notes": (
            "We'll use Anthropic's Claude Code extension for VS Code. Claude works through tools: "
            "it searches and reads to build its own picture of the code, and CLAUDE.md gives every "
            "session the project facts it can't infer. Edits and commands are tool calls too, and "
            "each one goes through a permission check you control. Plan mode separates thinking "
            "from doing, and rewind takes the code back to any earlier message."
        ),
    },
    {
        "kind": "Content",
        "layout": THREE_COL,
        "title": "Approvals and permission modes",
        "columns": [
            ("Manual", [
                "Edits and most commands ask",
                "For unfamiliar or risky work",
            ]),
            ("Edit automatically", [
                "Edits flow, most commands ask",
                "When tests check every change",
            ]),
            ("Plan", [
                "Reads and plans, changes nothing",
                "Before a change worth reviewing",
            ]),
        ],
        "notes": (
            "The three modes we'll use, picked from the indicator under the prompt box. There's "
            "a fourth, Auto, where a classifier decides which actions are safe, and recent "
            "versions start in it by default; set claudeCode.initialPermissionMode to manual to "
            "change that. Separate from the modes are allow and deny rules for specific "
            "commands. Don't ask again saves a rule just for you or for the whole team, in "
            "settings.local.json or settings.json. Demo 2 shows both."
        ),
    },
    {
        "kind": "Code",
        "layout": BASIC,
        "title": "A CLAUDE.md worth loading",
        "code": [
            "# Roombook",
            "## Commands",
            "- Run the tests: `uv run pytest`",
            "- Run the API: `uv run uvicorn roombook.main:app`",
            "## How it fits together",
            "- Times are stored in UTC, sent in office time",
            "- Every booking goes through policies.py",
            "## Rules",
            "- Don't edit a test to make code pass",
            "- Ask before adding a dependency",
        ],
        "notes": (
            "A trimmed version of the CLAUDE.md we'll build in Demo 2. Everything in it is loaded "
            "into every session, so it earns its place: the exact commands, the one convention "
            "that trips people up, and rules the code can't tell the agent. /init writes a good "
            "first draft, and then you edit it like a teammate's document. Generic advice like "
            "'write clean code' costs context and changes nothing."
        ),
    },
    {
        "kind": "Demo",
        "layout": BREAK,
        "title": "Demo 2: Context, tools, and approvals",
        "subtitle": "demos/02-claude-code · about 17 minutes",
        "notes": (
            "Switch to Manual. Read .claude/settings.json: the claudeMdExcludes patterns and the "
            "README deny. Trace POST /bookings, then /context. Send a subagent to survey the "
            "bookings table and run /context again: only its summary landed. /init, trim, add the "
            "rules. Coverage of policies.py: 71 percent. Generate tests, don't ask again for this "
            "project, Edit automatically, coverage near 100. Rewind, then settings.local.json."
        ),
    },
    {
        "kind": "Content",
        "layout": TWO_COL,
        "title": "Refactors and test generation",
        "columns": [
            ("Refactors", [
                "Plan first, and read it for behavior changes",
                "Tests pass before and after, unchanged",
                "Structure or behavior, never both at once",
            ]),
            ("Test generation", [
                "Expected values from real runs or the spec",
                "Pin the edges: limits, boundaries, the clock",
                "Coverage finds gaps; it isn't the goal",
            ]),
        ],
        "notes": (
            "Two high-leverage tasks. For a refactor, the plan is where you catch an agent that "
            "wants to tidy up behavior while it moves code, and unchanged tests passing is the "
            "proof. For tests of existing code, expected values come from running it; for new "
            "behavior, from the spec. Never from the agent's own arithmetic, which writes tests "
            "that agree with its mistakes. Coverage tells you what's missing, not what's safe."
        ),
    },
    {
        "kind": "Content",
        "layout": TWO_COL,
        "title": "Debugging and new features",
        "columns": [
            ("Debugging sessions", [
                "Reproduce with a failing test first",
                "Ask for the root cause and evidence",
                "Ask what else uses the faulty code",
            ]),
            ("New features", [
                "Start from a short written spec",
                "Review the plan before any code",
                "Build in slices, with tests in each",
            ]),
        ],
        "notes": (
            "Agents are eager to fix things, so slow them down. A failing test proves the bug and "
            "stays as a guard. Asking for the root cause with evidence, before any fix, stops the "
            "patch-the-symptom reflex, and asking what else uses the code finds the other places "
            "with the same fault. For features, a short spec beats a long chat, and the plan is "
            "the cheapest place to change direction. Demo 4 builds a feature this way."
        ),
    },
    {
        "kind": "Demo",
        "layout": BREAK,
        "title": "Demo 3: Debug and refactor with agents",
        "subtitle": "demos/03-debug-and-refactor · about 18 minutes",
        "notes": (
            "Read issue 42: an 8 pm New York booking shows on the next day. Claude in Manual: "
            "reproduce with a failing test, then the root cause, then what else uses day_bounds. "
            "Fix with tests for the room endpoint and the 25-hour day on 1 November. Review a wrong "
            "fix whose 48 tests pass. Then Copilot in Plan: extract the shared checks from "
            "bookings.py, read the Review Plan card for behavior changes, Implement Plan, tests "
            "unchanged and green."
        ),
    },
    {
        "kind": "Code",
        "layout": BASIC,
        "title": "A prompt with structure",
        "code": [
            "Goal:        Weekly recurring bookings, so teams",
            "             stop booking the same room by hand.",
            "Context:     Reuse check_slot in roombook/slots.py.",
            "Constraints: No new dependencies.",
            "             Don't change the existing tests.",
            "Done when:   A weekly 10:00 booking stays at 10:00",
            "             after New York's clocks change,",
            "             and every test passes.",
        ],
        "notes": (
            "Four parts: the goal and why, the context to reuse, the constraints that rule out "
            "shortcuts, and a definition of done the agent can check itself. This is a condensed "
            "version of the spec in Demo 4. The 'done when' line matters most, because an agent "
            "stops when it thinks it's done, and this tells it what done means, including the "
            "edge case most implementations get wrong."
        ),
    },
    {
        "kind": "Content",
        "layout": TWO_COL,
        "title": "When to step in, how to review",
        "columns": [
            ("Step in when the agent", [
                "Edits a test to make it pass",
                "Adds a dependency or touches unrelated files",
                "Repeats the same failing step",
            ]),
            ("Review the output by", [
                "Running the tests yourself",
                "Reading every diff, tests included",
                "Asking a second agent, then deciding",
            ]),
        ],
        "notes": (
            "You don't have to wait for the agent to finish. These three signals mean it has lost "
            "the thread, so stop it and say what you want instead. Reviewing is the part that "
            "doesn't change with agents: run the tests yourself, read the whole diff including the "
            "tests it wrote, and use a second agent as a reviewer with no stake in the first one's "
            "choices. Its findings are suggestions, and you decide."
        ),
    },
    {
        "kind": "Content",
        "layout": THREE_COL,
        "title": "Fitting agents into a team",
        "columns": [
            ("Repo hygiene", [
                "Clear commands to test and run",
                "Fast tests an agent can trust",
            ]),
            ("Context quality", [
                "One AGENTS.md for every tool",
                "Lessons learned go back into it",
            ]),
            ("Workflow fit", [
                "Shared rules and hooks in git",
                "Same review for agent changes",
            ]),
        ],
        "notes": (
            "What makes a repo easy for a new teammate makes it easy for an agent: obvious "
            "commands and fast, reliable tests. Keep one short, reviewed AGENTS.md that every tool "
            "reads, and feed lessons back into it. Check in shared permissions and hooks: a hook "
            "runs every time where an instruction is only a request, and CI is the backstop for "
            "what hooks don't cover. Agent work goes through normal review. Two more: agents read "
            "a lot, so usage adds up; keep tasks scoped. And they read untrusted text, so keep "
            "approvals on, deny risky commands, and keep secrets out."
        ),
    },
    {
        "kind": "Demo",
        "layout": BREAK,
        "title": "Demo 4: Ship a feature, keep it on track",
        "subtitle": "demos/04-team-feature · about 18 minutes",
        "notes": (
            "Read settings.json: allow, deny, and both hooks. Ask Claude to loosen a test and watch "
            "the hook block it. Move CLAUDE.md into AGENTS.md. Both agents plan the spec side by "
            "side; compare context, the clock change, constraints, and tests. Comment, send the "
            "feedback, approve Claude's plan. Stop the build on purpose for a test-first redirect. "
            "Review: tests, criterion 1, the diff, and a second opinion from Copilot's Ask agent."
        ),
    },
    {
        "kind": "Wrap",
        "layout": BASIC,
        "title": "Next steps",
        "bullets": [
            "Clone the repo and run the four demos",
            "Try one real task in each tool this week",
            "Write a short AGENTS.md for one of your repos",
            "Check in shared permissions and a formatting hook",
            "Start from the repo's starter kit and checklist",
        ],
        "notes": (
            "Point people to demos/README.md for setup. Suggest a small, real task with a clear "
            "finish line, such as a bug with a reproducible test, rather than a big feature. The "
            "starter kit in docs/starter-kit has an AGENTS.md template, team settings with both "
            "hooks, and a spec template. The checklist follows the same order as the session. "
            "Mention the Xebia Academy link in the README too."
        ),
    },
    {
        "kind": "Closing",
        "layout": END,
        "title": "Thank you",
        "subtitle": "Eric Greene · Xebia\neric.greene@xebia.com",
        "notes": (
            "Thank the audience. Leave this slide up while you answer any final questions, and "
            "repeat that the slides, demos, and checklist are in the repo."
        ),
    },
]


# ---------------------------------------------------------------------- validation

def words(text: str) -> int:
    """Count words the way a reader would: tokens containing a letter or digit."""
    return sum(1 for t in text.replace("`", "").split() if re.search(r"[A-Za-z0-9]", t))


def validate_data(slides: list[dict]) -> list[str]:
    problems = []
    if len(slides) > MAX_SLIDES:
        problems.append(f"{len(slides)} slides, limit is {MAX_SLIDES}")
    for n, s in enumerate(slides, start=1):
        groups = []
        if "bullets" in s:
            groups.append(s["bullets"])
        for _, items in s.get("columns", []):
            groups.append(items)
        total = sum(len(g) for g in groups)
        if total > MAX_BULLETS:
            problems.append(f"slide {n}: {total} bullets, limit is {MAX_BULLETS}")
        for g in groups:
            for b in g:
                if words(b) > MAX_WORDS:
                    problems.append(f"slide {n}: {words(b)} words in {b!r}")
        code_blocks = [s.get("code", [])] + [lines for _, lines in s.get("code_columns", [])]
        for block in code_blocks:
            if len(block) > MAX_CODE_LINES:
                problems.append(f"slide {n}: {len(block)} code lines, limit is {MAX_CODE_LINES}")
        if "image" in s and not (DIAGRAMS / s["image"]).is_file():
            problems.append(f"slide {n}: diagram {s['image']} is missing; export it first")
        if not s.get("notes", "").strip():
            problems.append(f"slide {n}: speaker notes are empty")
    for n in (1, len(slides)):
        if "eric.greene@xebia.com" not in slides[n - 1].get("subtitle", ""):
            problems.append(f"slide {n}: the presenter's email must be on the first and last slides")
    return problems


# ---------------------------------------------------------------------------- build

def add_runs(paragraph, text: str, size: int) -> None:
    """Split on backticks so `code` segments get the monospace font."""
    for i, part in enumerate(text.split("`")):
        if not part:
            continue
        run = paragraph.add_run()
        run.text = part
        run.font.size = Pt(size)
        if i % 2 == 1:
            run.font.name = MONO


def fill_lines(placeholder, lines: list[str], size: int) -> None:
    """Write one paragraph per line, without leaving a stray empty first paragraph."""
    tf = placeholder.text_frame
    tf.clear()
    first = True
    for line in lines:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        add_runs(para, line, size)


def fill_code(placeholder, lines: list[str], size: int) -> None:
    """Write code as unbulleted monospace lines.

    The body placeholder bullets every paragraph, which would put a dot in front of
    each line of code, so each paragraph gets an explicit buNone and no indent.
    """
    tf = placeholder.text_frame
    tf.clear()
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p_pr = para._p.get_or_add_pPr()
        p_pr.set("marL", "0")
        p_pr.set("indent", "0")
        p_pr.insert(0, p_pr.makeelement(qn("a:buNone"), {}))
        run = para.add_run()
        # A non-breaking space keeps blank lines and leading indentation intact.
        run.text = line.replace(" ", " ") if line else " "
        run.font.size = Pt(size)
        run.font.name = MONO


def transparent_png(path: Path) -> io.BytesIO:
    """Make the diagram's white page background transparent.

    The diagram tool exports on solid white, which would cover the template's light
    background artwork. Pillow comes with python-pptx, so no extra dependency is needed.
    """
    img = Image.open(path).convert("RGBA")
    img.putdata([(r, g, b, 0) if min(r, g, b) >= 250 else (r, g, b, a)
                 for r, g, b, a in img.get_flattened_data()])
    out = io.BytesIO()
    img.save(out, format="PNG")
    out.seek(0)
    return out


def add_diagram(slide, image: str, crop: tuple[float, float]) -> None:
    """Place a diagram under the title, scaled to the body area and cropped.

    The diagram tool always exports a full landscape page, so the empty margins above
    and below the drawing are cropped off before the picture is sized.
    """
    top_crop, bottom_crop = crop
    pic = slide.shapes.add_picture(transparent_png(DIAGRAMS / image), Inches(DIAGRAM_LEFT_IN),
                                   Inches(DIAGRAM_TOP_IN))
    pic.crop_top, pic.crop_bottom = top_crop, bottom_crop
    aspect = pic.image.size[0] / (pic.image.size[1] * (1 - top_crop - bottom_crop))
    height = min(DIAGRAM_MAX_HEIGHT_IN, DIAGRAM_MAX_WIDTH_IN / aspect)
    pic.height, pic.width = Inches(height), Inches(height * aspect)


def move(placeholder, *, top: float, height: float | None = None) -> None:
    """Reposition a placeholder vertically, keeping its inherited left and width.

    A slide placeholder inherits its geometry from the layout until it's given its own,
    and python-pptx needs all four values set to write one, so the current left and
    width are carried over explicitly.
    """
    left, width = placeholder.left, placeholder.width
    new_height = Inches(height) if height is not None else placeholder.height
    placeholder.left, placeholder.width = left, width
    placeholder.top, placeholder.height = Inches(top), new_height


def restore_slide_number(slide) -> None:
    """python-pptx skips slide-number placeholders when cloning a layout, but the
    template's own slides carry them, so copy the layout's one across."""
    for lp in slide.slide_layout.placeholders:
        if lp.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
            slide.shapes._spTree.append(copy.deepcopy(lp._element))


def remove_unfilled(slide, filled: set[int]) -> None:
    """An unfilled picture placeholder renders as an empty frame, so anything we
    didn't write to is removed rather than left behind."""
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx not in filled:
            ph._element.getparent().remove(ph._element)


def build(output: Path) -> None:
    prs = Presentation(TEMPLATE)
    layouts = prs.slide_masters[0].slide_layouts
    assert len(layouts) == 29, f"template has {len(layouts)} layouts, expected 29"

    # Remove the template's example slides. Layouts, the master and the notes master
    # survive, which is all the deck needs.
    rel_id = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    id_list = prs.slides._sldIdLst
    for sld in list(id_list):
        prs.part.drop_rel(sld.get(rel_id))
        id_list.remove(sld)
    assert len(prs.slides) == 0

    for s in SLIDES:
        slide = prs.slides.add_slide(layouts[s["layout"]])
        ph = {p.placeholder_format.idx: p for p in slide.placeholders}
        filled = {0}
        ph[0].text_frame.text = s["title"]

        if s["layout"] in (COVER, BREAK, END):
            fill_lines(ph[1], s["subtitle"].split("\n"), SUBTITLE_PT)
            filled.add(1)
        elif "columns" in s:
            slots = THREE_COL_SLOTS if s["layout"] == THREE_COL else TWO_COL_SLOTS
            for (head_idx, body_idx), (heading, items) in zip(slots, s["columns"]):
                fill_lines(ph[head_idx], [heading], COLUMN_HEADING_PT)
                fill_lines(ph[body_idx], items, COLUMN_BODY_PT)
                filled.update({head_idx, body_idx})
                if s["layout"] == TWO_COL:
                    move(ph[head_idx], top=TWO_COL_HEADING_TOP_IN)
                    move(ph[body_idx], top=TWO_COL_BODY_TOP_IN,
                         height=TWO_COL_BODY_HEIGHT_IN)
        elif "code_columns" in s:
            for (head_idx, body_idx), (heading, lines) in zip(TWO_COL_SLOTS, s["code_columns"]):
                fill_lines(ph[head_idx], [heading], COLUMN_HEADING_PT)
                fill_code(ph[body_idx], lines, COLUMN_CODE_PT)
                filled.update({head_idx, body_idx})
                move(ph[head_idx], top=TWO_COL_HEADING_TOP_IN)
                move(ph[body_idx], top=TWO_COL_BODY_TOP_IN, height=TWO_COL_BODY_HEIGHT_IN)
        elif "code" in s:
            fill_code(ph[1], s["code"], CODE_PT)
            filled.add(1)
        elif "image" in s:
            add_diagram(slide, s["image"], s["crop"])
        else:
            fill_lines(ph[1], s["bullets"], BULLET_PT)
            filled.add(1)

        remove_unfilled(slide, filled)
        if s["layout"] not in (COVER, END):
            restore_slide_number(slide)
        slide.notes_slide.notes_text_frame.text = s["notes"]

    assert len(prs.slides) == len(SLIDES)
    for n, slide in enumerate(prs.slides, start=1):
        assert slide.notes_slide.notes_text_frame.text.strip(), f"slide {n} has no notes"

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def inventory() -> str:
    rows = [f"{'#':>2}  {'kind':<8} title"]
    for n, s in enumerate(SLIDES, start=1):
        title = s["title"].replace("\n", " ")
        rows.append(f"{n:>2}  {s['kind']:<8} {title}")
    return "\n".join(rows)


def check_script() -> list[str]:
    """Compare the presenter script's slide headings with the deck, line for line.

    Each slide in PRESENTER_SCRIPT.md has a heading of the form "## Slide N: Title",
    and the titles drift whenever a slide is renamed without updating the script.
    """
    expected = [f"## Slide {n}: {s['title'].replace(chr(10), ' ')}"
                for n, s in enumerate(SLIDES, start=1)]
    text = SCRIPT.read_text(encoding="utf-8") if SCRIPT.exists() else ""
    found = [line.rstrip() for line in text.splitlines() if line.startswith("## Slide ")]
    problems = []
    for i in range(max(len(expected), len(found))):
        want = expected[i] if i < len(expected) else "(no slide)"
        got = found[i] if i < len(found) else "(missing)"
        if want != got:
            problems.append(f"expected {want!r}, found {got!r}")
    return problems


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--inventory", action="store_true",
                        help="print the slide list and exit, for checking the script against it")
    parser.add_argument("--check-script", action="store_true",
                        help="check PRESENTER_SCRIPT.md's slide headings against the deck and exit")
    args = parser.parse_args()

    if args.inventory:
        print(inventory())
        return 0

    if args.check_script:
        problems = check_script()
        for p in problems:
            print(f"  - {p}", file=sys.stderr)
        print("presenter script matches the deck" if not problems
              else f"{len(problems)} heading(s) don't match", file=sys.stderr if problems else sys.stdout)
        return 1 if problems else 0

    problems = validate_data(SLIDES)
    if problems:
        print("Slide data breaks the rules in CLAUDE.md:", file=sys.stderr)
        for p in problems:
            print(f"  - {p}", file=sys.stderr)
        return 1

    build(OUTPUT)
    print(f"built {OUTPUT.relative_to(ROOT)} with {len(SLIDES)} slides")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
